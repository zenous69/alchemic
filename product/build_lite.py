#!/usr/bin/env python3
"""
Build the FREE "lite" lead-magnet from the full kit.

It keeps the useful data entry (Facturas, Gastos, Plantilla factura) and a basic
quarterly IVA (Modelo 303) summary, but reserves the automatic Modelo 130 (IRPF),
the annual summary and the PDF guide for the paid version. A community-friendly
free asset that warms buyers and upsells to the full kit.

Output: product/Kit-Autonomo-Espana-GRATIS.xlsx
Run build_kit.py first (it needs the full workbook as input).
"""
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(__file__)
SRC = os.path.join(HERE, "Kit-Autonomo-Espana.xlsx")
OUT = os.path.join(HERE, "Kit-Autonomo-Espana-GRATIS.xlsx")

NAVY, ACCENT, WHITE = "1F3A5F", "F2C200", "FFFFFF"
UPSELL_URL = "https://TU-ENLACE-DE-VENTA"  # owner replaces with the Gumroad URL

wb = load_workbook(SRC)

# 1) drop paid-only sheet
if "Resumen anual" in wb.sheetnames:
    wb.remove(wb["Resumen anual"])

# 2) Resumen trimestral -> keep only IVA / Modelo 303 (cols A-D), drop 130 block
wr = wb["Resumen trimestral"]
for rng in list(wr.merged_cells.ranges):
    wr.unmerge_cells(str(rng))
wr.delete_cols(5, 9)               # remove columns E..M (the Modelo 130 block)
# clear the old note rows that referenced removed columns
for row in range(8, 30):
    for col in range(1, 6):
        wr.cell(row=row, column=col).value = None
wr.merge_cells("A1:D1")
wr["A1"] = "RESUMEN TRIMESTRAL  ·  IVA (Modelo 303)"

# upsell teaser box
wr.merge_cells("A9:D13")
box = wr["A9"]
box.value = ("🔒 VERSIÓN GRATUITA\n\n"
             "El cálculo AUTOMÁTICO del Modelo 130 (IRPF), el resumen anual y la "
             "Guía PDF del autónomo están en la versión completa.\n\n"
             f"Consíguela aquí:  {UPSELL_URL}")
box.font = Font(bold=True, color=NAVY, size=11)
box.fill = PatternFill("solid", fgColor=ACCENT)
box.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

# 3) mark Inicio as the free version + upsell line
wi = wb["Inicio"]
wi["B3"] = "VERSIÓN GRATUITA · controla facturas, gastos e IVA (Modelo 303)"
# find a free row under the intro to add the upsell
wi["B40"] = f"👉 Versión completa (Modelo 130 + resumen anual + Guía PDF): {UPSELL_URL}"
wi["B40"].font = Font(bold=True, color=NAVY, size=11)
wi.merge_cells("B40:E40")

wb.save(OUT)
print(f"OK -> {OUT}")
